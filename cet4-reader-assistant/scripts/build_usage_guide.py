from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
OUT_PATH = BASE_DIR / "CET4_阅读助手使用教程.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(88, 96, 105)
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"


def set_font(run, name: str = "Calibri", size: float | None = None, color: RGBColor | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_table_width(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_dxa):
        grid_col.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(24)
    title.font.color.rgb = INK
    title.font.bold = True
    title.paragraph_format.space_after = Pt(8)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(84)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CET-4 阅读助手")
    set_font(run, size=28, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("使用教程")
    set_font(run, size=16, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("截图识别、就地加词、导出生词，一套流程刷四级阅读")
    set_font(run, size=11, color=MUTED)

    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_margins(cell)
    set_cell_shading(cell, CALLOUT_FILL)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(
        "适合场景：把四级阅读截图变成可点击文本，边刷题边收集不会的单词。"
    )
    set_font(run, size=11, color=INK, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("版本日期：2026-05-17")
    set_font(run, size=10, color=MUTED)
    doc.add_page_break()


def add_shortcut_table(doc: Document) -> None:
    rows = [
        ("查看释义", "单击单词", "右侧显示中文、音标、词性"),
        ("加入生词", "双击单词", "最快的加词方式"),
        ("高亮/取消高亮", "选中文本后右键", "立即切换高亮状态"),
        ("搜索文章", "Ctrl + F", "聚焦搜索框"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [1800, 2100, 5460])
    headers = ("目的", "操作", "结果")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(text)
        set_font(run, size=10.5, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(text)
            set_font(run, size=10.5)
    set_table_width(table, [1800, 2100, 5460])


def add_steps(doc: Document) -> None:
    doc.add_heading("1. 第一次使用", level=1)
    steps = [
        "双击 run.bat。",
        "如果是第一次打开，程序会自动创建环境并安装依赖；安装完成后会自动打开主界面。",
        "如果电脑里默认 Python 版本不适合运行 OCR，程序会自动换用可用版本重新创建环境。",
        "如果电脑弹出网络或权限提示，允许安装流程继续执行。",
        "如果安装失败，启动窗口会停住显示错误，不会再一闪而过；看完提示后重新运行即可。",
        "首次安装成功后，以后通常直接双击 run.bat 即可。",
    ]
    for item in steps:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("2. 每天刷题的标准流程", level=1)
    steps = [
        "打开软件。",
        "按 Ctrl + Shift + A，框选四级阅读题里的正文区域。",
        "等待 OCR 完成，正文会自动显示在左侧。",
        "遇到不会的词，直接在文章里双击加入。",
        "需要看释义时，单击这个词，右侧就会显示中文、音标和词性。",
        "做完一篇后，点击“导出 vocab.txt”。",
    ]
    for item in steps:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("3. 什么图片最适合上传", level=1)
    doc.add_paragraph(
        "最适合的是正文页，或者只框住正文区域的截图。这样文章最干净，后续查词也最顺。"
    )
    for item in (
        "优先上传 Passage 正文页，不要优先上传只有题目和选项的页面。",
        "拍照时尽量让页面正、清楚、少反光，正文尽量占画面主体。",
        "如果一张照片里同时拍进另一页、题目区、页脚或反面透字，软件会尽量筛掉，但仍可能留下少量杂字。",
        "如果识别结果杂乱，重新截图时只框正文区域，通常比整页照片更稳。",
    ):
        doc.add_paragraph(item, style="List Bullet")


def add_usage_sections(doc: Document) -> None:
    doc.add_heading("4. 你会用到的几个区域", level=1)

    doc.add_heading("文章区", level=2)
    doc.add_paragraph(
        "左侧显示 OCR 后的完整文章。已经加入过的生词会被高亮，方便你回看。"
    )

    doc.add_heading("生词区", level=2)
    doc.add_paragraph(
        "右侧上半部分会显示当前保存过的单词，自动去重、自动转小写、自动按字母排序。"
    )

    doc.add_heading("查词区", level=2)
    doc.add_paragraph(
        "右侧下半部分显示所点单词的中文、音标、词性和来源。本地词典优先命中；若本地没有，程序会异步尝试在线接口。"
    )

    doc.add_heading("5. 快捷操作", level=1)
    add_shortcut_table(doc)

    doc.add_heading("6. 四级阅读专用清理", level=1)
    doc.add_paragraph("程序会自动尽量过滤这些干扰内容：")
    for item in (
        "Questions 46-50 这类题组标题",
        "Questions 51 to 55 are based on the following passage 这类说明行",
        "Section A / Section B",
        "Passage One / Passage Two",
        "A / B / C / D 选项行",
        "页码",
        "题号",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "如果截图里混进了很多题干、表格或边栏，建议重新框选得更贴近正文，识别效果会更稳。"
    )

    doc.add_heading("7. 导出和历史", level=1)
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [2700, 6660])
    for cell, text in zip(table.rows[0].cells, ("文件", "用途")):
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        run = cell.paragraphs[0].add_run(text)
        set_font(run, size=10.5, color=INK, bold=True)
    for file_name, meaning in (
        ("vocab/vocab.txt", "你手动导出的总生词表，一行一个词。"),
        ("vocab/history/YYYY-MM-DD_HH-MM-SS_vocab.txt", "每次导出时生成的历史快照，便于回看刷过的词。"),
    ):
        cells = table.add_row().cells
        for cell, text in zip(cells, (file_name, meaning)):
            set_cell_margins(cell)
            run = cell.paragraphs[0].add_run(text)
            set_font(run, size=10.5)
    set_table_width(table, [2700, 6660])


def add_troubleshooting(doc: Document) -> None:
    doc.add_heading("8. 遇到问题先看这里", level=1)
    rows = [
        ("双击后以前会闪退", "这是旧版安装脚本的问题；新版会显示错误并自动避开不兼容的 Python 版本。"),
        ("截图热键没反应", "点击工具栏里的“截图”按钮继续使用；有些电脑会拦截全局热键。"),
        ("识别出一堆题号和选项", "重新截图，只框正文区域，少带上下边缘。"),
        ("上传题目页后内容很乱", "题目页不是最合适的输入；优先上传正文页，或只截 Passage 正文。"),
        ("单词没有中文释义", "程序会先查本地词典，再尝试在线接口；若仍失败，可检查 data/dictionaries/ 里的离线词典文件。"),
        ("导出的单词顺序很乱", "不用手动整理，软件导出时会自动按字母排序。"),
        ("想删掉某个生词", "在右侧生词列表里选中后点“移除选中”，或双击列表中的词。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [2700, 6660])
    for cell, text in zip(table.rows[0].cells, ("情况", "处理方式")):
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        run = cell.paragraphs[0].add_run(text)
        set_font(run, size=10.5, color=INK, bold=True)
    for left, right in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (left, right)):
            set_cell_margins(cell)
            run = cell.paragraphs[0].add_run(text)
            set_font(run, size=10.5)
    set_table_width(table, [2700, 6660])

    doc.add_heading("9. 最推荐的实际用法", level=1)
    p = doc.add_paragraph()
    p.add_run("一篇文章一气呵成：").bold = True
    p.add_run(
        "先截图识别，再一路读完，遇到不会的词就双击收下。做完题后统一导出，再回看今天的生词表。这样最接近真正刷题，不会被频繁切窗口打断。"
    )


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.paragraphs[0].add_run("CET-4 阅读助手使用教程")
        set_font(run, size=9, color=MUTED)


def build() -> Path:
    doc = Document()
    set_styles(doc)
    add_cover(doc)
    add_steps(doc)
    add_usage_sections(doc)
    add_troubleshooting(doc)
    add_footer(doc)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
