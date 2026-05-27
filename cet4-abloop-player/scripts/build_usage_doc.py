from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT = PROJECT_ROOT / "docs" / "usage" / "CET4_听力精听播放器使用说明.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(91, 101, 112)
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "C9D3DF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_fixed_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths[idx] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_east_asia_font(run, font_name="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_run(paragraph, text, bold=False, color=None):
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def add_bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    add_run(para, text)
    return para


def add_numbered(doc, text):
    para = doc.add_paragraph(style="List Number")
    add_run(para, text)
    return para


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_table_width(table, [9360])
    set_table_borders(table, color="D7DEE8", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_run(p2, body)
    doc.add_paragraph()


def add_info_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_fixed_table_width(table, [1700, 7660])
    set_table_borders(table)
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "说明"
    for cell in hdr:
        set_cell_shading(cell, HEADER_FILL)
        for p in cell.paragraphs:
            for r in p.runs:
                set_east_asia_font(r)
                r.bold = True
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_east_asia_font(r)
    doc.add_paragraph()


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = DARK_BLUE
    title.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CET-4 听力精听播放器使用说明")
    set_east_asia_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("本地项目：D:\\Codex Project\\CET-4\\cet4-abloop-player")
    set_east_asia_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def build():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    add_run(title, "CET-4 听力精听播放器使用说明")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    add_run(subtitle, "适用材料：25-6-1 / cet4_2025_06_1.mp3、试题 PDF、解析 PDF", color=MUTED)

    add_callout(
        doc,
        "一句话怎么用",
        "打开本地网页，点左侧题号跳到对应听力位置；听不清时设置 A-B 循环反复听；校准时间点后保存 JSON，下次可继续用。",
    )

    doc.add_heading("1. 打开播放器", level=1)
    add_numbered(doc, "确认本地服务已经启动。本次我已经启动了服务，地址是 http://127.0.0.1:8000/cet4.html。")
    add_numbered(doc, "如果以后服务没开，在 PowerShell 输入下面两行命令。")
    add_callout(
        doc,
        "启动命令",
        'cd "D:\\Codex Project\\CET-4\\cet4-abloop-player"\n.\\.venv\\Scripts\\python.exe scripts\\local_server.py',
    )
    add_numbered(doc, "浏览器打开 http://127.0.0.1:8000/cet4.html。")
    add_bullet(doc, "如果默认音频没有自动加载，点右上角“导入 MP3”，选择 materials/listening/25-6-1 文件夹里的 cet4_2025_06_1.mp3。")

    doc.add_heading("2. 日常练习流程", level=1)
    add_numbered(doc, "在左侧点击 1-25 的题号，播放器会跳到该题的初始时间点。")
    add_numbered(doc, "按播放键开始听；中间区域可以改倍速，例如 0.75x、0.9x、1.25x。")
    add_numbered(doc, "右侧查看对应原文、答案和解析；原文里的 [1]、[2] 等标记也可以点击跳转到题目。")
    add_numbered(doc, "右侧“AI 时间轴”会列出本题附近的自动识别片段。点时间片段可以直接跳转播放。")
    add_numbered(doc, "如果某个片段想反复听，点这一行右侧的“循环”，它会自动设置 A-B 区间并开始循环。")
    add_numbered(doc, "听不清某一句时，用“设 A”和“设 B”圈出范围，再点“A-B 循环”。")
    add_numbered(doc, "如果只想循环本题，点“本题 A-B”，再点“A-B 循环”。")

    doc.add_heading("3. 使用 AI 时间轴精听", level=1)
    add_bullet(doc, "AI 时间轴由本地 faster-whisper base.en 模型生成，已经放进播放器。")
    add_bullet(doc, "每一行左侧是时间范围，例如 00:49-00:54；右侧是模型识别出的英文片段。")
    add_bullet(doc, "点整行会跳到该片段开头播放。")
    add_bullet(doc, "点“循环”会把该片段设为 A-B 循环，适合反复听某一句或半句。")
    add_bullet(doc, "AI 识别文字可能有个别词不准，背诵和核对时以“原文”区域为准。")

    doc.add_heading("4. 校准题目时间点", level=1)
    add_bullet(doc, "当前 1-25 题时间点是初始估算值，第一次使用建议边听边校准。")
    add_bullet(doc, "跳到某题后，拖动播放器到准确开始位置，点“当前时间设为本题”。")
    add_bullet(doc, "也可以直接在“开始 / 结束”输入框里填时间，例如 03:25 或 205。")
    add_bullet(doc, "调整完点“更新时间”。")
    add_bullet(doc, "全部或部分校准后，点右上角“保存 JSON”，导出自己的时间点文件。")
    add_bullet(doc, "下次打开网页后，点“导入 JSON”，选择上次保存的文件，就能恢复校准后的时间点。")

    doc.add_heading("5. 按钮速查", level=1)
    add_info_table(
        doc,
        [
            ("导入 MP3", "手动选择本地音频。默认路径正常时可以不点。"),
            ("导入 JSON", "导入上次保存的题目时间点。"),
            ("保存 JSON", "保存当前 1-25 题的开始/结束时间。"),
            ("倍速", "调整播放速度，适合精听慢放或复盘快听。"),
            ("当前时间设为本题", "把播放器当前位置写成本题开始时间。"),
            ("设 A / 设 B", "手动设置循环片段的起点和终点。"),
            ("本题 A-B", "用当前题目的开始/结束时间作为循环范围。"),
            ("A-B 循环", "开启或关闭反复播放 A-B 区间。"),
            ("AI 时间轴片段", "点击某个时间片段，直接跳到该片段播放。"),
            ("AI 时间轴循环", "点击片段右侧“循环”，自动把该片段设为 A-B 循环。"),
            ("搜索关键词", "搜索题干、选项、原文和解析。"),
        ],
    )

    doc.add_heading("6. 文件在哪里", level=1)
    add_info_table(
        doc,
        [
            ("播放器入口", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\cet4.html"),
            ("使用说明", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\docs\\usage\\CET4_听力精听播放器使用说明.docx"),
            ("题库数据", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\data\\cet4_2025_06_1.js"),
            ("AI 时间轴", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\data\\cet4_2025_06_1_timeline.js"),
            ("时间轴 JSON", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\data\\cet4_2025_06_1_timeline.json"),
            ("界面逻辑", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\js\\cet4.js"),
            ("界面样式", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\css\\cet4.css"),
            ("本地服务", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\scripts\\local_server.py"),
            ("生成时间轴脚本", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\scripts\\generate_timeline.py"),
            ("Whisper 模型", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\runtime\\models"),
            ("原始材料", "D:\\Codex Project\\CET-4\\cet4-abloop-player\\materials\\listening\\25-6-1"),
        ],
    )

    doc.add_heading("7. 常见问题", level=1)
    add_bullet(doc, "网页能打开但没声音：先确认浏览器没有静音，再点“导入 MP3”手动选择音频。")
    add_bullet(doc, "点击题号位置不准：这是正常的，初始时间是估算值；用“当前时间设为本题”校准后保存 JSON。")
    add_bullet(doc, "AI 时间轴文字和原文不完全一致：这是自动语音识别造成的，以标准原文为准；时间点仍可用来定位。")
    add_bullet(doc, "A-B 循环没生效：确认 A 和 B 都有值，并且 B 时间大于 A 时间。")
    add_bullet(doc, "下次打开时间点丢了：先导入上次保存的 JSON；浏览器本地缓存也会保存一份，但 JSON 更稳。")
    add_bullet(doc, "想停止服务：如果是在 PowerShell 前台启动的，按 Ctrl+C；如果是后台 pythonw 服务，可以在任务管理器结束 pythonw.exe。")

    doc.add_heading("8. 这次修改了什么", level=1)
    add_bullet(doc, "没有从零开发，保留了 ABLoopPlayer 上游源码。")
    add_bullet(doc, "新增四级专用页面，左侧题号、中间播放器、右侧原文答案解析。")
    add_bullet(doc, "新增 2025 年 6 月第一套听力数据，包含 1-25 题、答案、解析和原文。")
    add_bullet(doc, "新增 JSON 导入导出，方便保存自己校准过的题目时间点。")
    add_bullet(doc, "新增本地 faster-whisper base.en 模型生成的 210 个 AI 时间轴片段，并接入点击跳转和片段循环。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
